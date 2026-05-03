from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT_DIR = Path("data_modeling_updated/model_outputs")
REPORT_PATH = OUTPUT_DIR / "formatted_model_output_report.docx"

COMPARISON_LABELS = {
    "0": "Citation vs Warning",
    "1": "Arrest vs Warning",
}

TERM_LABELS = {
    "Intercept": "Intercept",
    "Black": "Black vs White",
    "Hispanic": "Hispanic vs White",
    "Asian/Pacific Islander": "Asian/Pacific Islander vs White",
    "Female": "Female vs Male",
    "wealth_c": "Centered WealthIndex (log scale)",
    "wealth_c:Black": "WealthIndex x Black",
    "wealth_c:Hispanic": "WealthIndex x Hispanic",
    "wealth_c:Asian/Pacific Islander": "WealthIndex x Asian/Pacific Islander",
    "wealth_c:Female": "WealthIndex x Female",
}

NUMERIC_COLUMNS = [
    "Coefficient",
    "Odds Ratio",
    "Cluster-Robust SE",
    "z-value",
]


def load_metric_table(file_name: str, metric_name: str) -> pd.DataFrame:
    table = pd.read_csv(OUTPUT_DIR / file_name, index_col=0)
    stacked = table.stack().rename(metric_name).reset_index()
    stacked.columns = ["term", "comparison", metric_name]
    stacked["comparison"] = stacked["comparison"].astype(str)
    return stacked


def prettify_term(term: str) -> str:
    if term in TERM_LABELS:
        return TERM_LABELS[term]
    if term.startswith("bs(age_c, df=4, include_intercept=False)["):
        basis_index = int(term.rsplit("[", 1)[1].rstrip("]")) + 1
        return f"Age spline basis {basis_index}"
    return term


def format_numeric(value: float) -> str:
    if pd.isna(value):
        return ""
    return f"{value:.3f}"


def format_pvalue(value: float) -> str:
    if pd.isna(value):
        return ""
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def significance_stars(pvalue: float) -> str:
    if pd.isna(pvalue):
        return ""
    if pvalue < 0.001:
        return "***"
    if pvalue < 0.01:
        return "**"
    if pvalue < 0.05:
        return "*"
    return ""


def build_combined_table() -> pd.DataFrame:
    metric_files = {
        "Coefficient": "mnlogit_coefficients.csv",
        "Odds Ratio": "mnlogit_odds_ratios.csv",
        "Cluster-Robust SE": "mnlogit_cluster_robust_se.csv",
        "z-value": "mnlogit_zvalues.csv",
        "p-value": "mnlogit_pvalues.csv",
    }

    merged: pd.DataFrame | None = None
    for metric_name, file_name in metric_files.items():
        metric_df = load_metric_table(file_name, metric_name)
        if merged is None:
            merged = metric_df
        else:
            merged = merged.merge(metric_df, on=["term", "comparison"], how="inner")

    if merged is None:
        raise ValueError("No model output tables could be loaded.")

    merged["Comparison"] = merged["comparison"].map(COMPARISON_LABELS)
    merged["Term"] = merged["term"].map(prettify_term)
    merged["Significance"] = merged["p-value"].apply(significance_stars)

    comparison_order = {label: i for i, label in enumerate(COMPARISON_LABELS.values())}
    merged["comparison_sort"] = merged["Comparison"].map(comparison_order)
    merged["term_sort"] = merged["term"]
    merged = merged.sort_values(["comparison_sort", "term_sort"]).drop(
        columns=["comparison_sort", "term_sort", "comparison", "term"]
    )

    ordered_columns = [
        "Comparison",
        "Term",
        "Coefficient",
        "Odds Ratio",
        "Cluster-Robust SE",
        "z-value",
        "p-value",
        "Significance",
    ]
    return merged[ordered_columns]


def format_table_for_output(table: pd.DataFrame) -> pd.DataFrame:
    formatted = table.copy()
    for column in NUMERIC_COLUMNS:
        formatted[column] = formatted[column].apply(format_numeric)
    formatted["p-value"] = formatted["p-value"].apply(format_pvalue)
    return formatted


def read_run_summary() -> list[str]:
    summary_path = OUTPUT_DIR / "run_summary.txt"
    return [line.strip() for line in summary_path.read_text().splitlines() if line.strip()]


def read_assumption_snapshot() -> dict:
    assumption_path = OUTPUT_DIR / "assumption_report.json"
    return json.loads(assumption_path.read_text())


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_cell_margins(cell, top: int = 80, start: int = 80, bottom: int = 80, end: int = 80) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    cell_margins = cell_properties.first_child_found_in("w:tcMar")
    if cell_margins is None:
        cell_margins = OxmlElement("w:tcMar")
        cell_properties.append(cell_margins)

    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        margin = cell_margins.find(qn(f"w:{key}"))
        if margin is None:
            margin = OxmlElement(f"w:{key}")
            cell_margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def format_table_cell(cell, bold: bool = False, font_size: int = 9) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_bullet_list(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        run.font.size = Pt(10)


def add_dataframe_table(document: Document, dataframe: pd.DataFrame, title: str) -> None:
    document.add_heading(title, level=2)
    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"
    table.autofit = True

    header_cells = table.rows[0].cells
    for index, column in enumerate(dataframe.columns):
        header_cells[index].text = str(column)
        format_table_cell(header_cells[index], bold=True, font_size=9)
        set_cell_shading(header_cells[index], "D9E2F3")

    for row_index, (_, row) in enumerate(dataframe.iterrows()):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cells[column_index].text = str(value)
            format_table_cell(cells[column_index], bold=False, font_size=9)
            if row_index % 2 == 1:
                set_cell_shading(cells[column_index], "F7F9FC")

    document.add_paragraph("")


def set_page_layout(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)


def build_docx_report() -> Document:
    combined = build_combined_table()
    formatted = format_table_for_output(combined)
    run_summary_lines = read_run_summary()
    assumptions = read_assumption_snapshot()

    missingness = assumptions["missingness"]
    class_balance = assumptions["class_balance"]
    recommendations = assumptions["recommendations"]
    max_missing_variable = max(missingness, key=missingness.get)
    max_missing_share = missingness[max_missing_variable]

    document = Document()
    set_page_layout(document)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Multinomial Logit Model Output Report")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Consolidated coefficients, odds ratios, cluster-robust standard errors, z-values, and p-values")
    subtitle_run.italic = True
    subtitle_run.font.name = "Times New Roman"
    subtitle_run.font.size = Pt(10)

    document.add_heading("Model Context", level=1)
    add_bullet_list(document, run_summary_lines)

    document.add_heading("How To Read This Document", level=1)
    add_bullet_list(
        document,
        [
            "Outcome reference category: Warning.",
            "Race reference category: White.",
            "Sex reference category: Male.",
            "Standard errors are cluster-robust.",
            "WealthIndex is already logged, so wealth terms are on the log-wealth scale.",
            "Age is modeled with spline basis terms; those coefficients are shown for completeness but should not be interpreted one-by-one as standalone substantive effects.",
        ],
    )

    document.add_heading("Consolidated Coefficient Tables", level=1)
    for comparison_label in COMPARISON_LABELS.values():
        comparison_df = formatted[formatted["Comparison"] == comparison_label].drop(columns=["Comparison"])
        add_dataframe_table(document, comparison_df, comparison_label)

    document.add_heading("Assumption Snapshot", level=1)
    add_bullet_list(
        document,
        [
            f"Largest missingness is in {max_missing_variable} at {max_missing_share:.2%}.",
            f"Outcome mix is imbalanced: Warning {class_balance['1']:.2%}, Citation {class_balance['2']:.2%}, Arrest {class_balance['3']:.2%}.",
            f"Zero-count race x sex cells: {assumptions['zero_count_race_sex_cells']}.",
            f"Maximum VIF: {assumptions['max_vif']:.3f}.",
            f"Age nonlinearity flag: {'Yes' if assumptions['age_nonlinearity_flag'] else 'No'}.",
            f"Wealth nonlinearity flag: {'Yes' if assumptions['wealth_nonlinearity_flag'] else 'No'}.",
            f"IIA sensitivity flag: {'Yes' if assumptions['iia_flag'] else 'No'}; wealth relative change = {assumptions['iia_relative_change']['wealth_c']:.3f}.",
            f"Independence note: {assumptions['independence_note']}",
        ],
    )

    document.add_heading("Recommendations Carried Forward", level=1)
    add_bullet_list(document, recommendations)

    document.add_heading("Significance Legend", level=1)
    add_bullet_list(
        document,
        [
            "* p < 0.05",
            "** p < 0.01",
            "*** p < 0.001",
        ],
    )

    return document


def main() -> None:
    document = build_docx_report()
    document.save(REPORT_PATH)
    print(f"Wrote report to {REPORT_PATH}")


if __name__ == "__main__":
    main()
